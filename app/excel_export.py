"""Excel 台账导出：生成可直接并入现有台账的 xlsx。

排版要点：
- 工作表名 ``<n>月``（不带前导零），六列表头：开始时间 | 部门 | 需求/故障 | 原因 | 处理方法 | 完成时间
- 分组模式：每天先写一行日期分组行（A 列整数日期序列号，B~F 留空），
  其后是该日数据行（A 列时间小数，数字格式 ``h:mm``）
- 完整模式：A 列直接写完整 datetime 序列号，数字格式 ``m/d/yyyy h:mm``
- 两种模式通用：冻结 ``A2``、A4 纵向、页边距 左右 0.7 / 上下 0.75 / 页眉页脚 0.3、
  宋体 11 号垂直居中、默认行高 13.5
- 日期序列号基准日 1899-12-30，只做纯日期与当日秒数运算，不经过本地时区换算
  （按 UTC 语义计算，避免因时区差一天）
- 内嵌照片：等比缩放到宽 ≤200、高 ≤100 像素，锚定在数据行 G 列，
  每行最多处理前 1 张 + 处理后 1 张
"""

from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .config import UPLOAD_DIR
from .db import get_db
from .utils import parse_dt, today_local

# ---------- 常量 ----------

# 排版模式
GROUPED = "grouped"  # 按日期分组（默认）
FULL = "full"  # 每行完整日期时间

# 表头六列
HEADERS = ("开始时间", "部门", "需求/故障", "原因", "处理方法", "完成时间")

# 两种排版模式各自的列宽（单位：字符宽度）
COLUMN_WIDTHS = {
    GROUPED: (10.5, 14, 26.875, 36.125, 35.125, 37),
    FULL: (17.5, 11.875, 26.5, 38.875, 47.875, 17.5),
}

# A / F 列数字格式
FMT_DATE = "m/d/yyyy"
FMT_TIME = "h:mm"
FMT_DATETIME = "m/d/yyyy h:mm"

# 单元格字体与行高：宋体 11 号、垂直居中、默认行高 13.5
FONT_NAME = "宋体"
FONT_SIZE = 11
DEFAULT_ROW_HEIGHT = 13.5

# 页面设置：A4、纵向、左右 0.7 / 上下 0.75 / 页眉页脚 0.3
PAPER_A4 = 9
PAGE_MARGINS = {
    "left": 0.7,
    "right": 0.7,
    "top": 0.75,
    "bottom": 0.75,
    "header": 0.3,
    "footer": 0.3,
}

# Excel 日期序列号基准日
EXCEL_EPOCH = date(1899, 12, 30)

# 图片缩放上限（像素）与锚定列（G 列）
PHOTO_MAX_WIDTH = 200
PHOTO_MAX_HEIGHT = 100
PHOTO_COLUMN = 7
PHOTO_STAGES = ("before", "after")  # 处理前 / 处理后
PHOTOS_PER_STAGE = 1  # 每个阶段最多内嵌张数

# SQLite 单条语句的占位符上限留出余量，按批查询
_ID_CHUNK = 900

COLUMN_LETTERS = tuple(get_column_letter(index) for index in range(1, len(HEADERS) + 1))


# ---------- 通用小工具 ----------


def _get(row, key, default=None):
    """兼容 sqlite3.Row 与 dict 的取值，缺失或为 None 时返回默认值。"""
    try:
        value = row[key]
    except (KeyError, IndexError, TypeError):
        return default
    return default if value is None else value


def _as_dt(value):
    """把数据库文本 / date / datetime 统一成 datetime，失败返回 None，秒以下截断。"""
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    parsed = parse_dt(value)
    return parsed.replace(microsecond=0) if parsed else None


def _as_date(value):
    """把数据库文本 / date / datetime 统一成 date，失败返回 None。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = parse_dt(value)
    return parsed.date() if parsed else None


def _time_fraction(moment):
    """当日时间占比，取值 0 <= v < 1。"""
    seconds = moment.hour * 3600 + moment.minute * 60 + moment.second
    return seconds / 86400


def _serial(moment):
    """Excel 日期序列号：基准日 1899-12-30，整数部分为天数、小数为当日占比。

    仅由 date 差值与当日秒数得出，与本地时区无关，不会差一天。
    """
    return (moment.date() - EXCEL_EPOCH).days + _time_fraction(moment)


def _clean_text(value):
    """单元格文本去掉首尾空白，空文本返回 None（真正留空）。"""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_mode(group_mode):
    """归一化排版模式，非法值回退为分组模式。"""
    mode = str(group_mode or "").strip().lower()
    return FULL if mode == FULL else GROUPED


def _int_or_none(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


# ---------- 工单归一化 ----------


def _record_start(row):
    """工单开始时间：优先 reported_at，缺失时退化为 occurred_on 当天 0 点。"""
    start = _as_dt(_get(row, "reported_at"))
    if start is not None:
        return start
    day = _as_date(_get(row, "occurred_on"))
    return datetime.combine(day, time.min) if day else None


def _prepare_rows(records, year, month, include_open):
    """筛选并归一化待导出工单，按开始时间升序返回。

    - 只保留开始时间落在目标年月的工单（导出页已按月查询，这里再兜底一次）
    - 未闭环工单（无完成时间）仅在 include_open 为真时保留
    - 开始时间完全无法解析的工单跳过
    """
    items = []
    for index, row in enumerate(records or []):
        start = _record_start(row)
        if start is None or start.year != year or start.month != month:
            continue
        end = _as_dt(_get(row, "finished_at"))
        if end is None and not include_open:
            continue
        items.append(
            {
                "id": _get(row, "id"),
                "order": index,
                "day": start.date(),
                "start": start,
                "end": end,
                "texts": (
                    _clean_text(_get(row, "department")),
                    _clean_text(_get(row, "symptom")),
                    _clean_text(_get(row, "cause")),
                    _clean_text(_get(row, "solution")),
                ),
            }
        )
    items.sort(key=lambda item: (item["start"], item["order"]))
    return items


# ---------- 样式与页面设置 ----------


def _style_cell(cell):
    """统一宋体 11 号、垂直居中、关闭自动换行。"""
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE)
    cell.alignment = Alignment(vertical="center", wrap_text=False)


def _write_header(ws):
    for column, text in enumerate(HEADERS, start=1):
        _style_cell(ws.cell(row=1, column=column, value=text))


def _apply_sheet_settings(ws, mode):
    """冻结首行、列宽、默认行高与页面设置。"""
    ws.freeze_panes = "A2"
    ws.sheet_format.defaultRowHeight = DEFAULT_ROW_HEIGHT
    ws.sheet_format.customHeight = True
    for letter, width in zip(COLUMN_LETTERS, COLUMN_WIDTHS[mode]):
        ws.column_dimensions[letter].width = width
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = PAPER_A4
    ws.page_margins = PageMargins(**PAGE_MARGINS)


# ---------- 写行 ----------


def _write_group_row(ws, row_idx, day):
    """日期分组行：A 列整数日期序列号（格式 m/d/yyyy），B~F 留空。"""
    cell = ws.cell(row=row_idx, column=1, value=int(_serial(datetime.combine(day, time.min))))
    cell.number_format = FMT_DATE
    _style_cell(cell)
    # 空单元格也套用宋体，保证整表字体一致
    for column in range(2, 7):
        _style_cell(ws.cell(row=row_idx, column=column))


def _write_data_row(ws, row_idx, item, mode):
    """数据行：A 列按模式写时间小数或完整序列号，B~E 文本，F 列完成时间。"""
    start = item["start"]
    if mode == FULL:
        a_cell = ws.cell(row=row_idx, column=1, value=_serial(start))
        a_cell.number_format = FMT_DATETIME
    else:
        a_cell = ws.cell(row=row_idx, column=1, value=_time_fraction(start))
        a_cell.number_format = FMT_TIME
    _style_cell(a_cell)

    for column, text in enumerate(item["texts"], start=2):
        _style_cell(ws.cell(row=row_idx, column=column, value=text))

    end = item["end"]
    if end is None:
        return  # 未闭环工单：完成时间留空
    f_cell = ws.cell(row=row_idx, column=len(HEADERS))
    if end.date() == start.date():
        # 与开始时间同天：只写时间小数，格式 h:mm
        f_cell.value = _time_fraction(end)
        f_cell.number_format = FMT_TIME
    else:
        # 跨天完成：写完整 datetime 序列号
        f_cell.value = _serial(end)
        f_cell.number_format = FMT_DATETIME
    _style_cell(f_cell)


# ---------- 照片 ----------


def _photo_path(filename):
    """把库中的相对路径解析为上传统目录下的真实文件，越界或不存在返回 None。"""
    if not filename:
        return None
    relative = Path(str(filename))
    if relative.is_absolute():
        return None
    root = UPLOAD_DIR.resolve()
    path = (root / relative).resolve()
    try:
        path.relative_to(root)
    except ValueError:
        return None
    return path if path.is_file() else None


def _scaled_image(path):
    """读入图片并等比缩放到宽 ≤200、高 ≤100 像素（小图不放大），失败返回 None。"""
    try:
        image = XLImage(str(path))
    except Exception:
        return None  # 依赖缺失或文件损坏时跳过，不阻断整月导出
    width = float(image.width or 0)
    height = float(image.height or 0)
    if width <= 0 or height <= 0:
        return None
    scale = min(1.0, PHOTO_MAX_WIDTH / width, PHOTO_MAX_HEIGHT / height)
    if scale < 1.0:
        image.width = round(width * scale, 2)
        image.height = round(height * scale, 2)
    return image


def _embed_row_photos(ws, row_idx, record_id, photos):
    """把该行处理前 / 处理后各至多 1 张照片锚定到 G 列，返回实际嵌入张数。"""
    entry = (photos or {}).get(record_id) or {}
    count = 0
    for stage in PHOTO_STAGES:
        names = list(entry.get(stage) or [])[:PHOTOS_PER_STAGE]
        for name in names:
            path = _photo_path(name)
            if path is None:
                continue
            image = _scaled_image(path)
            if image is None:
                continue
            image.anchor = f"{get_column_letter(PHOTO_COLUMN)}{row_idx}"
            ws.add_image(image)
            count += 1
    return count


# ---------- 对外接口 ----------


def month_photos(record_ids):
    """按工单 id 读取照片文件名。

    返回 ``{record_id: {"before": [...], "after": [...]}}``，文件名是相对
    ``data/uploads`` 的相对路径，按 sort_order 排序；没有照片的工单不出现在结果里。
    """
    ids = []
    seen = set()
    for value in record_ids or []:
        rid = _int_or_none(value)
        if rid is None or rid in seen:
            continue
        seen.add(rid)
        ids.append(rid)
    if not ids:
        return {}

    result = {}
    conn = get_db()
    for start in range(0, len(ids), _ID_CHUNK):
        chunk = ids[start : start + _ID_CHUNK]
        placeholders = ",".join("?" * len(chunk))
        sql = (
            "SELECT record_id, stage, filename FROM record_photos "
            f"WHERE record_id IN ({placeholders}) "
            "ORDER BY record_id, stage, sort_order, id"
        )
        for row in conn.execute(sql, chunk):
            entry = result.setdefault(
                int(row["record_id"]), {"before": [], "after": []}
            )
            stage = row["stage"] if row["stage"] in PHOTO_STAGES else PHOTO_STAGES[0]
            filename = _clean_text(row["filename"])
            if filename:
                entry[stage].append(filename)
    return result


def _build(records, photos, year, month, group_mode, include_open, embed_photos):
    """内部实现：返回 (Workbook, 数据行数, 嵌入图片数)。"""
    mode = _normalize_mode(group_mode)
    today = today_local()
    # 年月缺省时按当前月份处理，保证导出口径与工作表名始终有值
    month_num = _int_or_none(month) or today.month
    year_num = _int_or_none(year) or today.year
    items = _prepare_rows(records, year_num, month_num, include_open)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{month_num}月"
    _write_header(ws)
    _apply_sheet_settings(ws, mode)

    image_count = 0
    row_idx = 2
    current_day = None
    for item in items:
        if mode == GROUPED and item["day"] != current_day:
            current_day = item["day"]
            _write_group_row(ws, row_idx, current_day)
            row_idx += 1
        _write_data_row(ws, row_idx, item, mode)
        if embed_photos:
            image_count += _embed_row_photos(ws, row_idx, item["id"], photos)
        row_idx += 1

    return wb, len(items), image_count


def export_month(
    records,
    photos,
    year,
    month,
    out_path,
    *,
    group_mode=GROUPED,
    include_open=False,
    embed_photos=False,
):
    """生成并保存该月台账 xlsx，返回 ``{"rows": 数据行数, "photos": 嵌入图片数, "sheet": 工作表名}``。

    ``rows`` 只统计数据行（不含表头与日期分组行）；``photos`` 是实际嵌入的图片张数，
    未开启内嵌时为 0。
    """
    wb, rows, photo_count = _build(
        records, photos, year, month, group_mode, include_open, embed_photos
    )
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"rows": rows, "photos": photo_count, "sheet": wb.active.title}
